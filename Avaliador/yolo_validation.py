from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from ultralytics import YOLO
import yaml


def load_class_names(data_yaml: Path) -> list[str]:
    with data_yaml.open("r", encoding="utf-8") as data_file:
        names = (yaml.safe_load(data_file) or {}).get("names", {})
    if isinstance(names, list):
        return [str(name) for name in names]
    return [str(value) for _, value in sorted((int(key), value) for key, value in names.items())]


def export_confusion_matrix(matrix, labels: list[str], output_file: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Matriz de confusao"
    # O Ultralytics armazena a matriz como [classe predita, classe real].
    sheet.cell(row=1, column=1, value="Predito \\ Real").font = Font(bold=True)

    for index, label in enumerate(labels, start=2):
        sheet.cell(row=1, column=index, value=label).font = Font(bold=True)
        sheet.cell(row=index, column=1, value=label).font = Font(bold=True)

    for row_index, row in enumerate(matrix, start=2):
        for column_index, value in enumerate(row, start=2):
            sheet.cell(row=row_index, column=column_index, value=int(value))

    sheet.freeze_panes = "B2"
    sheet.column_dimensions["A"].width = max(18, *(len(label) + 2 for label in labels))
    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_file)


def main() -> None:
    parser = argparse.ArgumentParser(description="Executa a validacao YOLO e exporta a matriz para XLSX.")
    parser.add_argument("--data", required=True, type=Path)
    parser.add_argument("--model", required=True, type=Path)
    parser.add_argument("--project", required=True, type=Path)
    parser.add_argument("--confidence", type=float)
    parser.add_argument("--device")
    parser.add_argument("--save-json", action="store_true")
    args = parser.parse_args()

    validation_args = {
        "data": str(args.data),
        "split": "test",
        "plots": True,
        "project": str(args.project),
        "save_json": args.save_json,
    }
    if args.confidence is not None:
        validation_args["conf"] = args.confidence
    if args.device:
        validation_args["device"] = args.device

    metrics = YOLO(str(args.model)).val(**validation_args)
    matrix = metrics.confusion_matrix.matrix
    class_names = load_class_names(args.data)
    labels = [*class_names, "background"]
    export_confusion_matrix(matrix, labels, Path(metrics.save_dir) / "confusion_matrix.xlsx")
    print(f"Matriz de confusao XLSX salva em: {Path(metrics.save_dir) / 'confusion_matrix.xlsx'}")


if __name__ == "__main__":
    main()
